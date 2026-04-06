#!/usr/bin/env python3
"""
Parse SpaceDC / faculty schedule .xlsx (e.g. "25-26 Spring sem - Расписание.xlsx").

Layout (per sheet, except "Staff load"):
  - A row contains "Vilnius time"; the following row lists slot labels like "9:00 - 10:30".
  - Column C: calendar date (datetime); column D: day label or subgroup index; column E: group code.
  - Time slots start at column F (index 5); "30 min" / non-range cells are skipped.

Rows that describe the same activity (same date, time window, title, location) for different cohorts
are merged into one JSON object with a `groups` list.

If several cohort rows show the same cell text for one slot (after merge fill), `groups` lists those
cohorts only. If only one row has text, that row's cohort is used (no whole-day expansion).

Excel merged cells: values are copied into every row/column covered by the merge (openpyxl only stores
the value on the top-left cell). Joint sessions use that shared text on each covered row so `groups`
matches the merge, not every group on the sheet for that day.

If a start time appears inside the event text (e.g. "Consultation 14:30"), that time is used
as the start and end = start + 1.5 hours (overriding the column slot).

Dependencies: pip install -r scripts/requirements-schedule.txt
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from collections import defaultdict
from typing import Any, Iterable

from openpyxl import load_workbook

TIME_RANGE_RE = re.compile(
    r"^\s*(\d{1,2})[:.](\d{2})\s*-\s*(\d{1,2})[:.](\d{2})\s*$",
    re.IGNORECASE,
)
# Start time inside title (not followed by " - " range)
EMBEDDED_START_RE = re.compile(
    r"(?<!\d)(\d{1,2})[:.](\d{2})(?!\s*-\s*\d)",
)
ROOM_RE = re.compile(r"(?i)aud\.?\s*(\d{3,4})\b")
# 25HR-JA, 24LR-CS1, 23HO, etc.
GROUP_CODE_RE = re.compile(
    r"\b\d{2}(?:LR|HR)-[A-Za-z0-9]+\b|\b\d{2}HO\b",
    re.IGNORECASE,
)
WEEK_LABEL_RE = re.compile(r"^\s*(\d+)\s*Week\s*$", re.IGNORECASE)
# Timetable tabs named like 25LR, 24LR (evening / remote track); no aud in cell → treat as online.
LR_SHEET_NAME_RE = re.compile(r"^\s*\d{2}LR\s*$", re.IGNORECASE)
EMPTY_CELL = frozenset({None, ""})


def _norm_cell(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, str):
        s = v.replace("\xa0", " ").strip()
        return s if s else None
    return str(v).strip() or None


def _parse_time_range(cell: str) -> tuple[time, time] | None:
    m = TIME_RANGE_RE.match(cell.strip())
    if not m:
        return None
    h1, m1, h2, m2 = (int(m.group(i)) for i in range(1, 5))
    return time(h1, m1), time(h2, m2)


def _combine(d: date, t: time) -> datetime:
    return datetime.combine(d, t)


def _add_hours(dt: datetime, hours: float) -> datetime:
    return dt + timedelta(hours=hours)


def _first_embedded_start(text: str) -> time | None:
    for m in EMBEDDED_START_RE.finditer(text):
        h, mi = int(m.group(1)), int(m.group(2))
        if 0 <= h <= 23 and 0 <= mi <= 59:
            return time(h, mi)
    return None


def _rooms_in(text: str) -> list[str]:
    return [m.group(1) for m in ROOM_RE.finditer(text)]


def _groups_in(text: str) -> list[str]:
    return list(dict.fromkeys(m.group(0).upper() for m in GROUP_CODE_RE.finditer(text)))


def _title_without_room(text: str) -> str:
    t = ROOM_RE.sub("", text)
    t = GROUP_CODE_RE.sub("", t)
    t = re.sub(r"\s+", " ", t).strip(" ,;-")
    return t


def _norm_raw_cell_for_merge(raw: str) -> str:
    """Stable fingerprint so the same pasted cell text on different group rows merges."""
    lines = sorted(
        re.sub(r"\s+", " ", ln.strip()).casefold()
        for ln in re.split(r"[\r\n]+", raw)
        if ln.strip()
    )
    return " | ".join(lines)


def _merge_continuation_lines(cell_str: str) -> list[str]:
    """Join split cells: title on one line, room or '(online)' on the next."""
    raw_lines = [ln.strip() for ln in re.split(r"[\r\n]+", cell_str) if ln.strip()]
    merged: list[str] = []
    for ln in raw_lines:
        if not merged:
            merged.append(ln)
            continue
        is_online_only = ln.lower() in ("(online)", "online")
        is_room_only = bool(re.fullmatch(r"(?i)aud\.?\s*\d{3,4}\s*", ln))
        if is_online_only or is_room_only:
            merged[-1] = f"{merged[-1]} {ln}".strip()
        else:
            merged.append(ln)
    return merged


@dataclass
class ScheduleEntry:
    sheet: str
    week_label: str | None
    calendar_date: str  # ISO date
    day_label: str | None
    groups: list[str]
    """Cohort codes (column E) merged when the same slot/title/location repeat for multiple rows."""
    slot_label: str | None
    start: str  # ISO datetime (naive, Vilnius wall time as stored)
    end: str
    title: str
    rooms: list[str]
    """Rooms parsed from the event text (e.g. aud 341 -> "341")."""
    source_line: str
    raw_cell: str


def _export_time_only(schedule_ts: str) -> str:
    """Internal 'YYYY-MM-DD HH:MM' -> 'HH:MM' for JSON (date is in `date`)."""
    if " " not in schedule_ts:
        return schedule_ts
    _, time_part = schedule_ts.split(" ", 1)
    if len(time_part) >= 5 and time_part[2] == ":":
        return time_part[:5]
    return time_part


def _entry_location(entry: ScheduleEntry) -> str | None:
    """
    Single place value for JSON: \"online\", room number (digits from `aud …`), or None.
    Online is detected from the cell line (word \"online\") or from title suffix \"(online)\".
    Sheets named like 25LR with no room in the cell default to \"online\" (LR track).
    """
    if re.search(r"\bonline\b", entry.source_line, re.I) or "(online)" in entry.title.casefold():
        return "online"
    if entry.rooms:
        return entry.rooms[0]
    if LR_SHEET_NAME_RE.match(entry.sheet):
        return "online"
    return None


def entry_to_json_obj(entry: ScheduleEntry) -> dict[str, Any]:
    return {
        "title": entry.title,
        "date": entry.calendar_date,
        "start": _export_time_only(entry.start),
        "end": _export_time_only(entry.end),
        "groups": entry.groups,
        "location": _entry_location(entry),
    }


def _norm_title_key(title: str) -> str:
    return re.sub(r"\s+", " ", title.strip()).casefold()


def find_sessions_with_multiple_groups(
    entries: list[ScheduleEntry],
    *,
    require_non_empty_rooms: bool = False,
) -> list[dict[str, Any]]:
    """
    Entries that share the same sheet, instant (start/end), title, and location get one bucket.
    If two or more distinct cohort codes appear across rows in that bucket, report as one joint session.

    Use unmerged parse output (merge_shared=False) so each cohort row is still separate.
    """
    buckets: dict[tuple[str, str, str, str, str | None, str], list[ScheduleEntry]] = {}
    for e in entries:
        if require_non_empty_rooms and not e.rooms:
            continue
        loc = _entry_location(e)
        tkey = _norm_title_key(e.title)
        key = (e.sheet, e.calendar_date, e.start, e.end, loc, tkey)
        buckets.setdefault(key, []).append(e)

    out: list[dict[str, Any]] = []
    for key, bucket in buckets.items():
        sheet, cal, start, end, loc, _tkey = key
        all_groups = sorted({g.upper() for x in bucket for g in x.groups if g})
        if len(all_groups) < 2:
            continue
        titles = list(dict.fromkeys(x.title for x in bucket))
        out.append(
            {
                "sheet": sheet,
                "calendar_date": cal,
                "start": start,
                "end": end,
                "location": loc,
                "title": titles[0] if len(titles) == 1 else titles,
                "groups": all_groups,
                "occurrence_count": len(bucket),
            }
        )
    out.sort(key=lambda r: (r["calendar_date"], r["start"], r["sheet"], ",".join(r["groups"])))
    return out


def merge_shared_events(entries: list[ScheduleEntry]) -> list[ScheduleEntry]:
    """One record per distinct activity; union column-E groups when date/time/cell/line match."""
    buckets: dict[tuple[str, str, str, str, str, str], list[ScheduleEntry]] = {}
    for e in entries:
        key = (
            e.sheet,
            e.calendar_date,
            e.start,
            e.end,
            _norm_raw_cell_for_merge(e.raw_cell),
            _norm_raw_cell_for_merge(e.source_line),
        )
        buckets.setdefault(key, []).append(e)

    merged: list[ScheduleEntry] = []
    for group_entries in buckets.values():
        first = group_entries[0]
        row_groups = sorted({g.upper() for e in group_entries for g in e.groups if g})
        raw_cells = list(dict.fromkeys(e.raw_cell for e in group_entries))
        raw_cell = raw_cells[0] if len(raw_cells) == 1 else " | ".join(raw_cells)
        day = next((e.day_label for e in group_entries if e.day_label), None)
        week = next((e.week_label for e in group_entries if e.week_label), None)
        merged.append(
            ScheduleEntry(
                sheet=first.sheet,
                week_label=week,
                calendar_date=first.calendar_date,
                day_label=day,
                groups=row_groups,
                slot_label=first.slot_label,
                start=first.start,
                end=first.end,
                title=first.title,
                rooms=first.rooms,
                source_line=first.source_line,
                raw_cell=raw_cell,
            )
        )
    merged.sort(key=lambda e: (e.calendar_date, e.start, e.title.casefold(), ",".join(e.groups)))
    return merged


def _find_header_rows(rows: list[tuple[Any, ...]]) -> tuple[int, int] | None:
    """Return (header_row_idx, time_row_idx) in 0-based indexing."""
    for i, row in enumerate(rows):
        for c in row:
            if isinstance(c, str) and "Vilnius time" in c:
                if i + 1 < len(rows):
                    return i, i + 1
                return None
    return None


def _build_slot_columns(time_row: tuple[Any, ...]) -> list[tuple[int, str, time, time]]:
    """List of (col_index, raw_label, start, end)."""
    slots: list[tuple[int, str, time, time]] = []
    for col_idx, cell in enumerate(time_row):
        s = _norm_cell(cell)
        if not s:
            continue
        parsed = _parse_time_range(s)
        if not parsed:
            continue
        t0, t1 = parsed
        slots.append((col_idx, s, t0, t1))
    return slots


def _is_blank_excel_value(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, str):
        return not v.replace("\xa0", " ").strip()
    return False


def _fill_merged_cell_values(ws, grid: list[list[Any]]) -> None:
    """Copy top-left value of each merged range into the other cells (same as Excel display)."""
    for mr in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = mr.min_col, mr.min_row, mr.max_col, mr.max_row
        r0, c0 = min_row - 1, min_col - 1
        if r0 >= len(grid):
            continue
        if c0 >= len(grid[r0]):
            continue
        val = grid[r0][c0]
        if _is_blank_excel_value(val):
            continue
        for r in range(min_row - 1, max_row):
            if r >= len(grid):
                break
            row = grid[r]
            for c in range(min_col - 1, max_col):
                if c >= len(row):
                    continue
                if r == r0 and c == c0:
                    continue
                if _is_blank_excel_value(row[c]):
                    row[c] = val


def _sheet_rows(ws) -> list[tuple[Any, ...]]:
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    ranges = getattr(ws.merged_cells, "ranges", None)
    if ranges:
        _fill_merged_cell_values(ws, grid)
    return [tuple(r) for r in grid]


def _slot_fill_index(
    rows: list[tuple[Any, ...]],
    time_row_idx: int,
    slot_col_indices: list[int],
) -> dict[tuple[date, int], list[tuple[str, str]]]:
    """For each (calendar day, slot column): cohort + cell text for every non-empty cell."""
    slot_fill: dict[tuple[date, int], list[tuple[str, str]]] = defaultdict(list)
    cur_date: date | None = None

    for row in rows[time_row_idx + 1 :]:
        c = row[2] if len(row) > 2 else None
        new_d: date | None = None
        if isinstance(c, datetime):
            new_d = c.date()
        elif isinstance(c, date):
            new_d = c
        if new_d is not None:
            cur_date = new_d
        if cur_date is None:
            continue
        g = _norm_cell(row[4] if len(row) > 4 else None)
        if not g:
            continue
        for col_idx in slot_col_indices:
            cell = _norm_cell(row[col_idx] if col_idx < len(row) else None)
            if cell:
                slot_fill[(cur_date, col_idx)].append((g.upper(), cell))
    return slot_fill


def _emission_groups(
    calendar_date: date,
    col_idx: int,
    row_group: str,
    line_codes: list[str],
    slot_fill: dict[tuple[date, int], list[tuple[str, str]]],
) -> list[str]:
    """
    Union cohorts that share identical slot text (vertical merge → same text on multiple rows).
    Single filled row → that row's cohort only (not every group in the day's block).
    """
    rg = row_group.upper()
    from_line = {x.upper() for x in line_codes}
    fills = slot_fill.get((calendar_date, col_idx), [])
    if len(fills) <= 1:
        return sorted({rg} | from_line)
    texts = {_norm_raw_cell_for_merge(t) for _, t in fills}
    if len(texts) == 1:
        return sorted({g for g, _ in fills} | from_line)
    return sorted({rg} | from_line)


def parse_sheet(
    sheet_name: str,
    rows: list[tuple[Any, ...]],
    *,
    merge_shared: bool = True,
    expand_day_cohort_slots: bool = True,
) -> list[ScheduleEntry]:
    hdr = _find_header_rows(rows)
    if not hdr:
        return []
    _, time_row_idx = hdr
    time_row = rows[time_row_idx]
    slot_cols = _build_slot_columns(time_row)
    if not slot_cols:
        return []

    slot_indices = [c[0] for c in slot_cols]
    slot_fill: dict[tuple[date, int], list[tuple[str, str]]] = defaultdict(list)
    if expand_day_cohort_slots:
        slot_fill = _slot_fill_index(rows, time_row_idx, slot_indices)

    out: list[ScheduleEntry] = []
    current_week: str | None = None
    current_date: date | None = None

    for row in rows[time_row_idx + 1 :]:
        b = row[1] if len(row) > 1 else None
        c = row[2] if len(row) > 2 else None
        d = row[3] if len(row) > 3 else None
        e = row[4] if len(row) > 4 else None

        bs = _norm_cell(b)
        if bs and WEEK_LABEL_RE.match(bs):
            current_week = bs.strip()

        if isinstance(c, datetime):
            current_date = c.date()
        elif isinstance(c, date):
            current_date = c

        if current_date is None:
            continue

        group_cell = _norm_cell(e)
        day_label = _norm_cell(d) if isinstance(d, str) else None
        if isinstance(d, int | float) and not isinstance(d, bool):
            day_label = None
        # subgroup rows: D is 2,3,... and E is the group code
        if group_cell is None:
            continue

        group = group_cell

        for col_idx, slot_label, slot_start, slot_end in slot_cols:
            if col_idx >= len(row):
                continue
            raw = row[col_idx]
            cell_str = _norm_cell(raw)
            if not cell_str:
                continue

            for line in _merge_continuation_lines(cell_str):
                if not line:
                    continue
                emb = _first_embedded_start(line)
                if emb:
                    start_dt = _combine(current_date, emb)
                    end_dt = _add_hours(start_dt, 1.5)
                else:
                    start_dt = _combine(current_date, slot_start)
                    end_dt = _combine(current_date, slot_end)

                rooms = _rooms_in(line)
                line_groups = _groups_in(line)
                if expand_day_cohort_slots:
                    groups_union = _emission_groups(
                        current_date,
                        col_idx,
                        group,
                        line_groups,
                        slot_fill,
                    )
                else:
                    groups_union = sorted({group.upper(), *[x.upper() for x in line_groups]})
                title = _title_without_room(line)
                if re.search(r"\bonline\b", line, re.I):
                    title = (title + " (online)").strip() if "(online)" not in title.lower() else title

                out.append(
                    ScheduleEntry(
                        sheet=sheet_name,
                        week_label=current_week,
                        calendar_date=current_date.isoformat(),
                        day_label=day_label,
                        groups=groups_union,
                        slot_label=slot_label,
                        start=start_dt.isoformat(sep=" ", timespec="minutes"),
                        end=end_dt.isoformat(sep=" ", timespec="minutes"),
                        title=title or line,
                        rooms=rooms,
                        source_line=line,
                        raw_cell=cell_str,
                    )
                )

    return merge_shared_events(out) if merge_shared else out


def parse_workbook(
    path: Path,
    sheets: Iterable[str] | None = None,
    *,
    merge_shared: bool = True,
    expand_day_cohort_slots: bool = True,
) -> list[ScheduleEntry]:
    # read_only=False: merged cell ranges are required to replicate values into covered rows.
    wb = load_workbook(path, read_only=False, data_only=True)
    names = list(sheets) if sheets else [n for n in wb.sheetnames if n.strip().lower() != "staff load"]
    all_entries: list[ScheduleEntry] = []
    for name in names:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        rows = _sheet_rows(ws)
        all_entries.extend(
            parse_sheet(
                name,
                rows,
                merge_shared=merge_shared,
                expand_day_cohort_slots=expand_day_cohort_slots,
            )
        )
    wb.close()
    return all_entries


def main() -> None:
    p = argparse.ArgumentParser(description="Parse schedule xlsx to JSON.")
    p.add_argument(
        "xlsx",
        type=Path,
        nargs="?",
        default=Path(__file__).resolve().parent.parent / "25-26 Spring sem - Расписание.xlsx",
        help="Path to the .xlsx file (default: repo root sibling of scripts/)",
    )
    p.add_argument(
        "--sheet",
        action="append",
        dest="sheets",
        help="Only parse these sheet names (repeatable). Default: all except 'Staff load'.",
    )
    p.add_argument("-o", "--output", type=Path, help="Write JSON to this file (UTF-8).")
    p.add_argument("--indent", type=int, default=2, help="JSON indent (default 2).")
    p.add_argument(
        "--list-multi-group",
        action="store_true",
        help="Print joint sessions: same time+title+location with 2+ cohort rows (uses unmerged parse).",
    )
    p.add_argument(
        "--multi-group-rooms-only",
        action="store_true",
        help="With --list-multi-group, only buckets that have at least one parsed room (aud …).",
    )
    p.add_argument(
        "--strict-row-groups",
        action="store_true",
        help="Do not use cross-row slot matching; each row emits only its column-E cohort (merge_shared still unions identical cells).",
    )
    args = p.parse_args()

    if not args.xlsx.is_file():
        raise SystemExit(f"File not found: {args.xlsx}")

    if args.list_multi_group:
        raw = parse_workbook(
            args.xlsx,
            args.sheets,
            merge_shared=False,
            expand_day_cohort_slots=not args.strict_row_groups,
        )
        joints = find_sessions_with_multiple_groups(
            raw,
            require_non_empty_rooms=args.multi_group_rooms_only,
        )
        text = json.dumps(joints, ensure_ascii=False, indent=args.indent if args.indent > 0 else None)
        print(text)
        if not joints:
            print(
                "Note: no joint sessions found. This mode only detects when two or more cohort rows "
                "have the same date, start, end, title, and room(s) — i.e. the same activity spelled "
                "the same way on multiple group lines. In this workbook, shared lessons are usually "
                "entered on one cohort row only; the other rows are empty for that slot, so nothing matches.\n"
                "The main JSON (-o schedule.json) uses `groups`: you still get one code per row unless "
                "the same cell text is repeated on multiple rows.",
                file=sys.stderr,
            )
        return

    entries = parse_workbook(
        args.xlsx,
        args.sheets,
        expand_day_cohort_slots=not args.strict_row_groups,
    )
    payload = [entry_to_json_obj(e) for e in entries]
    text = json.dumps(payload, ensure_ascii=False, indent=args.indent if args.indent > 0 else None)
    if args.output:
        args.output.write_text(text + "\n", encoding="utf-8")
    else:
        print(text)


if __name__ == "__main__":
    main()
